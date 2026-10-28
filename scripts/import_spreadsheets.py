"""One-time import of BlueCredit's existing spreadsheets into BlueCredit Pulse.

Usage:
    python scripts/import_spreadsheets.py /path/to/folder/with/files [--reset]

Any file that isn't in the folder is skipped. Re-running without --reset adds
duplicates for opportunities/funding, so use --reset for a clean reload.
Contacts are de-duplicated by email.
"""
import csv
import json
import os
import re
import sys
from datetime import date, datetime

import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from app import db as dbmod  # noqa: E402
from app.schema import field_names  # noqa: E402

USER = "import"
stats = {}


# ------------------------------------------------------------------ helpers
def s(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def to_date(v, dayfirst=False):
    if v in (None, ""):
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    t = s(v)
    for fmt in (["%d/%m/%Y"] if dayfirst else []) + ["%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%m/%d/%Y", "%m/%d/%y",
                                                     "%B %d, %Y", "%b %d, %Y", "%B %d %Y", "%d %B %Y"]:
        try:
            return datetime.strptime(t, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    m = re.search(r"(\d{1,2})/(\d{1,2})/(\d{2,4})", t)  # e.g. "next deadline is 7/27/25"
    if m:
        mo, d, y = map(int, m.groups())
        y = y + 2000 if y < 100 else y
        try:
            return date(y, mo, d).isoformat()
        except ValueError:
            return ""
    return ""


def to_num(v):
    if v in (None, ""):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    t = s(v).lower().replace(",", "")
    if re.search(r"\d\s*-\s*\$?\d", t):  # ranges stay as text
        return None
    m = re.search(r"\$?\s*(\d+(?:\.\d+)?)\s*(k|m|million|b)?", t)
    if not m:
        return None
    n = float(m.group(1))
    mult = {"k": 1e3, "m": 1e6, "million": 1e6, "b": 1e9}.get(m.group(2) or "", 1)
    return n * mult


def owner(v):
    t = s(v)
    for o in ("Barasa", "Sohaib", "Robert", "Beckee"):
        if o.lower() in t.lower():
            return t if "," in t else o
    return t


def sheet_rows(path, sheet=None, header_row=1):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    for _ in range(header_row - 1):
        next(rows, None)
    header = [s(h) for h in next(rows)]
    for r in rows:
        if r is None or not any(c not in (None, "") for c in r):
            continue
        yield {header[i]: r[i] for i in range(min(len(header), len(r))) if header[i]}


def csv_rows(path):
    with open(path, encoding="utf-8-sig", errors="replace") as f:
        yield from csv.DictReader(f)


class Loader:
    def __init__(self, conn):
        self.db = conn
        self.emails = {}
        self.name_org = set()
        for r in conn.execute("SELECT id, email, name, organization FROM contacts"):
            for e in split_emails(r["email"]):
                self.emails[e] = r["id"]
            self.name_org.add((s(r["name"]).lower(), s(r["organization"]).lower()))

    def add(self, entity, data, meta=None, src=""):
        allowed = set(field_names(entity))
        data = {k: v for k, v in data.items() if k in allowed and v not in (None, "")}
        if not data:
            return None
        ts = dbmod.now_iso()
        cols = list(data) + ["meta", "created_at", "updated_at", "updated_by"]
        vals = list(data.values()) + [json.dumps(meta, default=str) if meta else None, ts, ts, USER]
        cur = self.db.execute(f'INSERT INTO {entity} ({",".join(chr(34)+c+chr(34) for c in cols)}) '
                              f'VALUES ({",".join("?"*len(cols))})', vals)
        stats[f"{entity} ← {src}"] = stats.get(f"{entity} ← {src}", 0) + 1
        return cur.lastrowid

    def contact(self, data, meta=None, src=""):
        data["source"] = data.get("source") or src
        emails = split_emails(data.get("email"))
        for e in emails:
            if e in self.emails:  # fill blanks on the existing record, note the extra source
                rid = self.emails[e]
                row = dict(self.db.execute("SELECT * FROM contacts WHERE id=?", (rid,)).fetchone())
                upd = {k: v for k, v in data.items() if v and not row.get(k) and k in field_names("contacts")}
                tags = [t for t in s(row.get("tags")).split("; ") if t]
                if src and src not in tags and src != row.get("source"):
                    upd["tags"] = "; ".join(tags + [src])
                if upd:
                    self.db.execute(f'UPDATE contacts SET {",".join(chr(34)+k+chr(34)+"=?" for k in upd)} WHERE id=?',
                                    (*upd.values(), rid))
                stats["contacts merged (same email)"] = stats.get("contacts merged (same email)", 0) + 1
                return rid
        if not emails:
            key = (s(data.get("name")).lower(), s(data.get("organization")).lower())
            if key == ("", ""):
                return None
            if key in self.name_org:
                return None
            self.name_org.add(key)
        rid = self.add("contacts", data, meta, src)
        for e in emails:
            self.emails[e] = rid
        return rid


def split_emails(v):
    return [e.lower() for e in re.findall(r"[\w.+'-]+@[\w-]+(?:\.[\w-]+)+", s(v))]


FUND_STATUS = {"not started": "Identified", "": "Identified", "missed": "Missed", "rejected": "Lost",
               "out of scope": "Out of Scope", "submitted": "Submitted", "withdrawn": "Withdrawn",
               "in progress": "Pursuing", "won": "Won", "awarded": "Won"}
OPP_STAGE = {"not started": "Identified", "": "Identified", "in progress": "Pursuing", "submitted": "Submitted",
             "missed": "Missed", "rejected": "Lost", "not qualified": "Not Qualified", "cancelled": "Cancelled",
             "won": "Won", "awarded": "Won"}
FUND_TYPE = {"grant": "Grant", "accelerator": "Accelerator", "pitch comp": "Pitch Competition",
             "pitch competition": "Pitch Competition", "fellowship": "Fellowship", "vc": "VC", "loan": "Loan",
             "award": "Award", "incubator": "Incubator", "resources": "Resources", "seed round": "Seed Round",
             "sbir": "SBIR/STTR", "sttr": "SBIR/STTR"}


def fund_type(v):
    t = s(v).lower()
    return FUND_TYPE.get(t) or next((FUND_TYPE[k] for k in FUND_TYPE if k in t), "Other" if t else "")


# ------------------------------------------------------------------ importers
def tracker(L, path):
    for r in sheet_rows(path, "Funding Applications"):
        if not s(r.get("Application Name")):
            continue
        amt = r.get("Amount Requested")
        L.add("funding", {
            "name": s(r.get("Application Name")), "organization": s(r.get("Organization")),
            "type": fund_type(r.get("Type")) or fund_type(r.get("Type (raw)")),
            "status": FUND_STATUS.get(s(r.get("Status")).lower(), "Identified"),
            "deadline": to_date(r.get("Submission Deadline")), "amount": to_num(amt),
            "amount_text": s(amt), "pwin": to_num(r.get("PWin Score")), "owner": owner(r.get("Owner")),
            "contact_email": s(r.get("Contact Email")), "submitted_date": to_date(r.get("Submitted Date")),
            "decision_date": to_date(r.get("Decision Date")), "url": s(r.get("Application URL")),
            "notes": s(r.get("Notes")), "source": "BD Tracker"}, src="BD Tracker: Funding Applications")

    for r in sheet_rows(path, "Feasibility Study Opportunities"):
        if s(r.get("Program")):
            L.add("resources", {"name": s(r.get("Program")), "agency": s(r.get("Agency")),
                                "category": "Feasibility Program",
                                "notes": s(r.get("Typical Feasibility Studies / Analyses"))},
                  src="BD Tracker: Feasibility")

    for r in sheet_rows(path, "Consulting Opportunities"):
        if not s(r.get("Project Name")):
            continue
        L.add("opportunities", {
            "title": s(r.get("Project Name")), "kind": "Consulting", "agency": s(r.get("Client")),
            "stage": OPP_STAGE.get(s(r.get("Property")).lower(), "Identified"),
            "deadline": to_date(r.get("Submission Deadline")), "contact_email": s(r.get("Contact Email")),
            "owner": owner(r.get("Owner")), "pwin": to_num(r.get("PWin Score")), "priority": s(r.get("Priority")),
            "value": to_num(r.get("Project Value")), "url": s(r.get("URL")), "source": "BD Tracker"},
            src="BD Tracker: Consulting")

    for r in sheet_rows(path, "Subcontract Applications"):
        if not s(r.get("Company")):
            continue
        # NOTE: the "Username/Pw" column is deliberately NOT imported. Keep credentials in a password manager.
        notes = "\n".join(x for x in [s(r.get("Notes")), s(r.get("Contract Type")) and f"Contract type: {s(r.get('Contract Type'))}",
                                      s(r.get("Portal / registration link")) and f"Vendor portal: {s(r.get('Portal / registration link'))}"] if x)
        L.contact({"organization": s(r.get("Company")), "email": s(r.get("Contact email")),
                   "phone": s(r.get("Contact phone")), "category": "Prime Contractor",
                   "status": "Contacted" if s(r.get("Next step")) or s(r.get("Last touch")) else "To Contact",
                   "last_touch": to_date(r.get("Last touch")), "next_step": s(r.get("Next step")),
                   "owner": owner(r.get("Owner")), "notes": notes, "tags": "Subcontracting"},
                  src="BD Tracker: Subcontract")

    for r in sheet_rows(path, "Corporate Water Outreach"):
        name, li = s(r.get("Contact")), s(r.get("Linkedin"))
        if name.startswith("http"):
            li = li or name
            name = re.sub(r"[-_]", " ", name.rstrip("/").split("/")[-1]).title()
        st = s(r.get("Response status")).lower()
        L.contact({"name": name, "organization": s(r.get("Company")), "email": s(r.get("Email")),
                   "role": s(r.get("Role")), "linkedin": li, "category": "Corporate",
                   "status": {"to contact": "To Contact", "contacted": "Contacted", "responded": "Responded"}.get(st, "To Contact"),
                   "notes": s(r.get("Follow-up notes")), "tags": "Corporate water outreach"},
                  src="BD Tracker: Corporate Outreach")

    for r in sheet_rows(path, "Water Credit Projects"):
        if not s(r.get("Project Name")):
            continue
        notes = "\n".join(f"{k}: {s(r.get(k))}" for k in ("Available Water Credits", "Expected Timeline",
                          "Project Category", "Project Stage", "Watershed") if s(r.get(k)))
        L.add("opportunities", {
            "title": s(r.get("Project Name")), "kind": "Water Credit Project",
            "agency": s(r.get("Project Developers")), "location": s(r.get("Watershed")),
            "stage": OPP_STAGE.get(s(r.get("Status")).lower(), "Identified"),
            "value": to_num(r.get("Project Value (USD)")), "notes": notes, "source": "BD Tracker"},
            meta={k: s(v) for k, v in r.items()}, src="BD Tracker: Water Credit Projects")


def epa(L, path):
    for r in sheet_rows(path, header_row=3):
        if not s(r.get("Record Number")):
            continue
        tier = s(r.get("Tier"))
        L.add("opportunities", {
            "title": s(r.get("Description"))[:160], "kind": "Forecast", "agency": "EPA " + s(r.get("Procurement Office")),
            "ref_no": s(r.get("Record Number")), "naics": s(r.get("NAICS Code")), "value_text": s(r.get("Estimated Dollars")),
            "stage": "Identified", "priority": "High" if "1" in tier[:7] else "Medium" if "2" in tier[:7] else "Low",
            "set_aside": s(r.get("Procurement Method")), "location": s(r.get("Place Of Performance")),
            "contact_name": s(r.get("EPA Contact")), "contact_email": s(r.get("Contact Email")),
            "notes": "\n".join(x for x in [tier, s(r.get("Description")),
                                           f"Target award: FY{s(r.get('Target Award Year'))} {s(r.get('Target Award Quarter'))}",
                                           s(r.get("Priority Notes"))] if x),
            "source": "EPA forecast"}, src="EPA Opportunities")
        if s(r.get("Contact Email")):
            L.contact({"name": s(r.get("EPA Contact")), "organization": "EPA", "email": s(r.get("Contact Email")),
                       "role": "Contracting POC", "category": "Government"}, src="EPA Opportunities")


def forecast(L, path):
    seen = set()
    for r in csv_rows(path):
        lid = s(r.get("Listing ID"))
        if not lid or lid in seen:
            continue
        seen.add(lid)
        notes = [s(r.get("Description"))]
        for k in ("Requirement Status", "Acquisition Phase", "Contract Type", "Extent Competed", "Award Type",
                  "Estimated Award FY-QTR", "Period of Performance", "Additional Information"):
            if s(r.get(k)) and s(r.get(k)) not in ("To Be Determined", "N/A"):
                notes.append(f"{k}: {s(r.get(k))}")
        sb = s(r.get("Content: Small Business Specialist Info (Name)"))
        if sb:
            notes.append(f"Small business specialist: {sb} {s(r.get('Content: Small Business Specialist Info (Email)'))}")
        L.add("opportunities", {
            "title": s(r.get("Title")), "kind": "Forecast",
            "agency": " / ".join(x for x in [s(r.get("Agency")), s(r.get("Organization"))] if x),
            "ref_no": lid, "naics": s(r.get("NAICS Code")), "value_text": s(r.get("Estimated Contract Value")),
            "stage": "Identified", "deadline": to_date(r.get("Estimated Solicitation Date")),
            "set_aside": s(r.get("Set Aside Type")), "location": s(r.get("Place of Performance State")),
            "contact_name": s(r.get("Content: Point of Contact (Name) For")),
            "contact_email": s(r.get("Point of Contact (Email)")), "notes": "\n".join(n for n in notes if n),
            "source": "Acquisition Gateway forecast"}, src="Forecast export")
        for nm, em, role in [(r.get("Content: Point of Contact (Name) For"), r.get("Point of Contact (Email)"), "Contracting POC"),
                             (r.get("Content: Small Business Specialist Info (Name)"),
                              r.get("Content: Small Business Specialist Info (Email)"), "Small Business Specialist")]:
            if s(em):
                L.contact({"name": s(nm).title() if s(nm).isupper() else s(nm), "email": s(em), "role": role,
                           "organization": s(r.get("Organization")) or s(r.get("Agency")), "category": "Government"},
                          src="Forecast export")


def bidboard(L, path):
    for r in sheet_rows(path, "Agency Bid Boards", header_row=3):
        if not s(r.get("Agency")):
            continue
        L.add("resources", {"name": f"{s(r.get('Agency'))} — {s(r.get('Agency Forecast Portal')) or 'bid board'}",
                            "category": "Forecast Portal", "agency": s(r.get("Agency")),
                            "url": s(r.get("Agency Forecast URL")),
                            "contact": " ".join(x for x in [s(r.get("Office / OSDBU Contact")), s(r.get("Email")), s(r.get("Phone"))] if x),
                            "notes": "\n".join(f"{k}: {s(r.get(k))}" for k in ("Product Area", "Simplified Acq. Notes",
                                                                                 "Below-SAT Strategy") if s(r.get(k)))},
              src="Fed Agency Bid Board")
        if s(r.get("Email")):
            L.contact({"name": s(r.get("Office / OSDBU Contact")), "organization": s(r.get("Agency")),
                       "email": s(r.get("Email")), "phone": s(r.get("Phone")), "role": "Small business office",
                       "category": "Government"}, src="Fed Agency Bid Board")
    for r in sheet_rows(path, "Universal Platforms", header_row=2):
        if s(r.get("Platform")):
            L.add("resources", {"name": s(r.get("Platform")), "category": "Bid Board", "url": s(r.get("URL")),
                                "notes": s(r.get("Best For"))}, src="Fed Agency Bid Board")
    for r in sheet_rows(path, "Below-SAT Strategy", header_row=2):
        if s(r.get("Agency")):
            L.add("resources", {"name": f"{s(r.get('Agency'))} below-SAT approach", "category": "Below-SAT Strategy",
                                "agency": s(r.get("Agency")), "contact": s(r.get("Key Contact")),
                                "notes": s(r.get("Below-SAT Action"))}, src="Fed Agency Bid Board")
    for r in sheet_rows(path, "District & Regional Officers", header_row=3):
        if s(r.get("Agency")):
            L.add("resources", {"name": f"{s(r.get('Agency'))} {s(r.get('District / Office'))}".strip(),
                                "category": "District Office", "agency": s(r.get("Agency")),
                                "url": s(r.get("Small Business Portal URL")), "contact": s(r.get("Phone")),
                                "notes": "\n".join(x for x in [s(r.get("Division / Region")), s(r.get("Geography Covered")),
                                                               s(r.get("Notes"))] if x)}, src="Fed Agency Bid Board")


def deq(L, path):
    firm = ""
    for r in sheet_rows(path):
        firm = s(r.get("Firm")) or firm
        L.contact({"name": s(r.get("Participant")), "organization": firm, "email": s(r.get("Email")),
                   "category": "Government" if "deq" in firm.lower() or "virginia" in firm.lower() else "Partner",
                   "tags": "DEQ contacts"}, src="DEQ contacts")


def mailing(L, path):
    for r in sheet_rows(path):
        L.contact({"name": " ".join(x for x in [s(r.get("First Name")), s(r.get("Last Name"))] if x),
                   "email": s(r.get("Email")), "category": "Mailing List"}, src="Mailing list")


def funding_db(L, path):
    today = date.today().isoformat()
    for r in sheet_rows(path):
        if not s(r.get("Program")):
            continue
        dl_raw = r.get("Deadline")
        dl = to_date(dl_raw)
        notes = s(r.get("Description"))
        if s(dl_raw) and not dl:
            notes = (notes + f"\nDeadline: {s(dl_raw)}").strip()
        L.add("funding", {"name": s(r.get("Program")), "type": fund_type(r.get("Type")), "status": "Catalog",
                          "region": s(r.get("State/Country")), "deadline": dl, "amount": to_num(r.get("Amount")),
                          "amount_text": s(r.get("Amount")), "url": s(r.get("Website")), "notes": notes,
                          "source": "Funding DB"},
              meta={"date_posted": to_date(r.get("Date Posted")), "open": bool(dl and dl >= today)},
              src="Funding DB (catalog)")


def sbic(L, path):
    for r in csv_rows(path):
        detail = (f"Manager: {s(r.get('Manager'))}\nVintage: {s(r.get('Vintage Year'))}  Fund size: ${to_num(r.get('Fund Size')) or 0:,.0f}"
                  f"  Avg investment: ${to_num(r.get('Average Investment')) or 0:,.0f}\n"
                  f"Strategy: {s(r.get('Investment Strategy'))} ({s(r.get('Fund Style'))})\n"
                  f"Making new investments: {s(r.get('Making New Investments?'))}")
        L.contact({"name": s(r.get("Investor Relations Name")), "organization": s(r.get("Name")),
                   "email": s(r.get("Investor Relations Email")), "phone": s(r.get("Investor Relations Phone")),
                   "role": "Investor relations", "category": "Investor",
                   "location": ", ".join(x for x in [s(r.get("City")), s(r.get("State"))] if x),
                   "status": "To Contact" if s(r.get("Making New Investments?")).lower() == "yes" else "Dormant",
                   "tags": "SBIC" + ("; investing" if s(r.get("Making New Investments?")).lower() == "yes" else ""),
                   "notes": detail}, meta=dict(r), src="SBIC list")


def simplified(L, path):
    with open(path, encoding="utf-8-sig", errors="replace") as f:
        rows = list(csv.reader(f))
    for r in rows[1:]:
        if len(r) < 3:
            continue
        for e in split_emails(r[2]):
            L.contact({"name": e.split("@")[0].replace(".", " ").title(), "email": e, "organization": s(r[1]),
                       "role": "Contracting officer (past simplified acquisition)", "category": "Government",
                       "notes": f"PIID {s(r[0])}"}, src="Simplified acquisitions")


def industrial(L, path):
    for r in sheet_rows(path, "AWS Certified Sites"):
        if not s(r.get("Account Name")):
            continue
        L.contact({"organization": s(r.get("Account Name")), "name": s(r.get("Certificate Name")),
                   "location": s(r.get("Site Country")), "category": "Corporate Prospect",
                   "tags": "AWS-certified site",
                   "notes": "\n".join(f"{k}: {s(r.get(k))}" for k in ("AWS Site Registration No", "Site Address", "Catchment",
                                      "Level Attained", "Valid Until", "Industry Sector") if s(r.get(k)))},
                  src="Industrial sites (AWS certified)")


def watercontacts(L, path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    names = wb.sheetnames
    def rows(sheet):
        it = wb[sheet].iter_rows(values_only=True)
        h = [s(x) for x in next(it)]
        for r in it:
            if r and any(c not in (None, "") for c in r):
                yield h, r
    if "companies" in names:
        for h, r in rows("companies"):
            d = dict(zip(h, r))
            if s(d.get("Company Name")):
                L.contact({"organization": s(d.get("Company Name")), "category": "Corporate Prospect",
                           "location": s(d.get("Location")), "linkedin": "",
                           "tags": "; ".join(x for x in [s(d.get("Type")), s(d.get("Sector"))] if x),
                           "notes": "\n".join(f"{k}: {s(d.get(k))}" for k in ("Link", "Goal", "Year", "Volume", "Projects") if s(d.get(k)))},
                          src="Water contacts: companies")
    if "ngo" in names:
        for h, r in rows("ngo"):
            d = dict(zip(h, r))
            if s(d.get("Partner")):
                L.contact({"organization": s(d.get("Partner")), "email": s(d.get("Email")), "category": "NGO",
                           "location": ", ".join(x for x in [s(d.get("City")), s(d.get("Country"))] if x),
                           "notes": "\n".join(x for x in [s(d.get("Website")), s(d.get("Category")), s(d.get("Address"))] if x)},
                          src="Water contacts: NGOs")
    if "agencies" in names:
        for h, r in rows("agencies"):
            d = dict(zip(h, r))
            if s(d.get("Name")):
                typ = s(d.get("Type"))
                L.contact({"organization": s(d.get("Name")), "category": "Utility" if re.search(r"water|utilit|sewer|authority", typ, re.I) else "Government",
                           "location": s(d.get("Primary Service Area")), "tags": typ, "notes": s(d.get("Website"))},
                          src="Water contacts: agencies")
    if "experts" in names:
        for _, r in rows("experts"):
            r = list(r) + [None] * 9
            name = " ".join(x for x in [s(r[0]), s(r[1])] if x and len(x) > 1)
            if s(r[2]).startswith("http"):  # LinkedIn-style layout
                L.contact({"name": name, "linkedin": s(r[2]), "organization": s(r[4]), "role": s(r[5]),
                           "category": "Expert"}, src="Water contacts: experts")
            else:
                L.contact({"name": name, "organization": s(r[2]), "role": s(r[3]), "tags": s(r[4]),
                           "location": s(r[5]) if s(r[5]) != "N/A" else "", "email": s(r[6]),
                           "phone": " / ".join(x for x in [s(r[7]), s(r[8])] if x), "category": "Expert"},
                          src="Water contacts: experts")
    if "dotsubcontractingdir" in names:
        for h, r in rows("dotsubcontractingdir"):
            r = list(r) + [None] * 5
            m = re.match(r"(.*?)\s*(\(?\d{3}\)?[\s.-]?\d{3}[\s.-]?\d{4}.*)?$", s(r[4]))
            L.contact({"organization": s(r[0]), "name": (m.group(1) if m else s(r[4])).title(),
                       "phone": (m.group(2) or "") if m else "", "location": s(r[1]), "category": "Prime Contractor",
                       "tags": "DOT subcontracting directory", "notes": f"NAICS {s(r[2])}: {s(r[3])}"},
                      src="Water contacts: DOT subcontracting")
    if "gsacontracts" in names:
        for h, r in rows("gsacontracts"):
            d = dict(zip(h, r))
            L.contact({"organization": s(d.get("Legal Business Name")), "category": "Prime Contractor",
                       "location": ", ".join(x for x in [s(d.get("Vendor Address City")), s(d.get("Vendor Address State"))] if x),
                       "tags": "GSA awardee",
                       "notes": "\n".join(f"{k}: {s(v)}" for k, v in d.items() if s(v) and k not in ("Legal Business Name",))},
                      src="Water contacts: GSA awards")
    stats["(skipped reference sheets)"] = ", ".join(n for n in names if n in ("uscounties", "us_uni+colle", "usprivateschools"))


def dbr(L, path):
    import docx
    doc = docx.Document(path)
    reviews, cur = [], None
    for p in doc.paragraphs:
        if p.style.name.startswith("Heading 1"):
            m = re.search(r"(\d{1,2})/(\d{1,2})/(\d{4})", p.text)
            cur = {"date": date(int(m.group(3)), int(m.group(1)), int(m.group(2))).isoformat() if m else "",
                   "title": p.text.strip(), "lines": [], "attendees": ""}
            reviews.append(cur)
        elif cur is not None:
            t = p.text.rstrip()
            if p.style.name.startswith("Heading 3") and "&" in t and not cur["attendees"]:
                cur["attendees"] = t
            elif p.style.name.startswith("Heading"):
                cur["lines"].append(f"\n## {t}")
            else:
                cur["lines"].append(t)
    for rv in reviews:
        notes = re.sub(r"\n{3,}", "\n\n", "\n".join(rv["lines"])).strip()
        ts = dbmod.now_iso()
        c = L.db.execute("""INSERT INTO reviews (cadence, review_date, title, attendees, notes, snapshot, created_by,
                         created_at, updated_at) VALUES ('daily',?,?,?,?,NULL,?,?,?)""",
                         (rv["date"], rv["title"], rv["attendees"], notes, USER, ts, ts))
        stats["reviews ← Daily Business Review doc"] = stats.get("reviews ← Daily Business Review doc", 0) + 1
        for line in rv["lines"]:  # "Task - Owner Name" lines become action items
            m = re.match(r"^(.{8,}?)\s*-\s*(Robert Wafula|Barasa Wangusi|Sohaib Anwar|Beckee|Rebecca Mwombe)\s*$", line.strip())
            if m:
                L.add("actions", {"text": m.group(1).strip(), "owner": owner(m.group(2)), "status": "Open",
                                  "review_id": c.lastrowid, "notes": f"From {rv['title']}"}, src="DBR action items")


FILES = [
    ("BlueCredit_Business_Development_Tracker.xlsx", tracker),
    ("BlueCredit_EPA_Opportunities.xlsx", epa),
    ("forecast_export.csv", forecast),
    ("FedAgency_BidBoard.xlsx", bidboard),
    ("deqcontacts.xlsx", deq),
    ("Daily_Business_Review_.docx", dbr),
    ("funding_db.xlsx", funding_db),
    ("sbic_contacts.csv", sbic),
    ("simplifiedacquistions.csv", simplified),
    ("industrialsites_db.xlsx", industrial),
    ("watercontacts_db.xlsx", watercontacts),
    ("BlueCredit_MailingList.xlsx", mailing),
]


def find(folder, name):
    norm = lambda x: re.sub(r"[^a-z0-9]", "", x.lower().rsplit(".", 1)[0])
    for f in os.listdir(folder):
        if norm(f) == norm(name) or norm(f).endswith(norm(name)):
            return os.path.join(folder, f)
    return None


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    folder = sys.argv[1]
    if "--reset" in sys.argv and os.path.exists(dbmod.DB_PATH):
        for ext in ("", "-wal", "-shm"):
            if os.path.exists(dbmod.DB_PATH + ext):
                os.remove(dbmod.DB_PATH + ext)
    dbmod.init_db()
    conn = dbmod.connect()
    L = Loader(conn)
    for name, fn in FILES:
        path = find(folder, name)
        if not path:
            print(f"  – {name}: not found, skipped")
            continue
        print(f"  • {name}")
        fn(L, path)
        conn.commit()
    dbmod.log(conn, USER, "system", None, "import", "Initial import from spreadsheets")
    conn.commit()
    print("\nImported:")
    for k, v in stats.items():
        print(f"  {k}: {v}")
    for t in ("opportunities", "funding", "contacts", "resources", "actions", "reviews"):
        print(f"  TOTAL {t}: {conn.execute(f'SELECT COUNT(*) FROM {t}').fetchone()[0]}")
    conn.close()


if __name__ == "__main__":
    main()
